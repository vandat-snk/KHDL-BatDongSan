"""Verify current saved data and provenance; no assumed test/visual/report claims."""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import sqlite3
import statistics
import sys

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'src'))
from artifact_contract import file_sha256,read_verified_source,validate_core_outputs,validate_derivative

def require(condition,message):
    if not condition: raise ValueError(message)

def close(left,right):
    return math.isclose(float(left),float(right),rel_tol=1e-12,abs_tol=1e-7)

def csv_rows(path):
    with path.open(encoding='utf-8-sig',newline='') as stream: return list(csv.DictReader(stream))

def test_evidence(root,run_id):
    path=root/'outputs/python_test_results.json'
    if not path.exists(): return {'status':'not_run_for_current_revision'}
    record=json.loads(path.read_text(encoding='utf-8'))
    current={p.relative_to(root).as_posix():file_sha256(p) for p in sorted((root/'tests').glob('test_*.py'))}
    if record.get('core_run_id')!=run_id or record.get('test_files')!=current:
        return {'status':'stale_test_record'}
    return {k:record[k] for k in ('status','tests_run','failures','errors','skipped')}


def dashboard_test_evidence(root,run_id):
    path=root/'outputs/dashboard_test_results.json'
    if not path.exists(): return {'status':'not_run_for_current_revision'}
    record=json.loads(path.read_text(encoding='utf-8'))
    if (record.get('core_run_id')!=run_id or
        record.get('test_file_sha256')!=file_sha256(root/'tests/test_dashboard.mjs') or
        record.get('dashboard_sha256')!=file_sha256(root/'outputs/DASHBOARD_BAT_DONG_SAN.html')):
        return {'status':'stale_test_record'}
    return {key:record[key] for key in ('status','scope','findings','exported_search_rows','clustered_rows')}

def verify(root=ROOT,require_workbook=False,core_only=False):
    manifest=validate_core_outputs(root)
    analysis=json.loads((root/'outputs/analysis.json').read_text(encoding='utf-8'))
    records=json.loads((root/'data/processed/listings_clean.json').read_text(encoding='utf-8'))
    raw,raw_hash=read_verified_source(root/'data/raw/chotot_data.json')
    saved_csv=csv_rows(root/'data/processed/listings_clean.csv')
    require(len(records)==len(saved_csv)==analysis['counts']['clean_count'],'CSV/JSON/counts disagree.')
    for record,row in zip(records,saved_csv):
        for key,expected in record.items():
            value=row[key]
            same=str(value).lower()==str(expected).lower() if isinstance(expected,bool) else close(value,expected) if isinstance(expected,(int,float)) else value==expected
            require(same,f'CSV mismatch {record["row_id"]}:{key}')
        require(record['area_m2']>0 and record['sale_price_vnd']>=10_000_000,'Invalid retained numeric domain.')
        require(close(record['price_per_m2_vnd']*record['area_m2'],record['sale_price_vnd']),'Unit-price arithmetic mismatch.')
    lineage=csv_rows(root/'data/processed/row_lineage.csv')
    rejected=csv_rows(root/'data/processed/listings_rejected.csv')
    require(len(lineage)==len(raw),'Lineage must cover every raw record.')
    require({int(x['row_id']) for x in lineage}==set(range(1,len(raw)+1)),'Missing/duplicate lineage IDs.')
    for status,key in [('kept','clean_count'),('duplicate','duplicate_count'),('rejected','rejected_count')]:
        require(sum(x['status']==status for x in lineage)==analysis['counts'][key],'Lineage counts mismatch.')
    require({r['row_id'] for r in records}=={int(x['row_id']) for x in lineage if x['status']=='kept'},'Kept identity mismatch.')
    require({int(r['row_id']) for r in rejected}=={int(x['row_id']) for x in lineage if x['status']=='rejected'},'Quarantine identity mismatch.')
    from pipeline import CORE_FIELDS,normalize,clean_records
    for row in lineage:
        if row['status']=='duplicate':
            left,right=raw[int(row['row_id'])-1],raw[int(row['canonical_row_id'])-1]
            require([normalize(left.get(k)) for k in CORE_FIELDS]==[normalize(right.get(k)) for k in CORE_FIELDS],'Duplicate content mismatch.')
    rerun,_,_,_=clean_records(raw)
    require(rerun.row_id.tolist()==[r['row_id'] for r in records],'Saved retained IDs differ from current cleaning code.')
    for key in ('area_m2','sale_price_vnd','price_per_m2_vnd'):
        require(all(close(x,r[key]) for x,r in zip(rerun[key],records)),f'Current numeric policy mismatch: {key}')
    for key in ('property_type','transaction_status','province','district','quality_flags'):
        require(rerun[key].tolist()==[r[key] for r in records],f'Current text policy mismatch: {key}')
    for key,field in [('median_area_m2','area_m2'),('median_price_vnd','sale_price_vnd'),('median_price_per_m2_vnd','price_per_m2_vnd')]:
        require(close(statistics.median(r[field] for r in records),analysis['overall'][key]),'Saved median mismatch.')
    benchmark=analysis['sql_benchmark']
    with sqlite3.connect((root/'outputs/real_estate.sqlite').as_uri()+'?mode=ro',uri=True) as connection:
        require(connection.execute('PRAGMA integrity_check').fetchone()[0]=='ok','Database integrity failed.')
        require(connection.execute('PRAGMA foreign_key_check').fetchall()==[],'Foreign-key failures.')
        rows=connection.execute('SELECT l.row_id,p.name,t.name,l.district,l.transaction_status,l.area_m2,l.sale_price_vnd,l.price_per_m2_vnd,l.is_outlier_group,l.cluster FROM listings l JOIN provinces p USING(province_id) JOIN property_types t USING(type_id) ORDER BY l.row_id').fetchall()
        require(len(rows)==len(records),'Database count mismatch.')
        keys=['row_id','province','property_type','district','transaction_status','area_m2','sale_price_vnd','price_per_m2_vnd','is_outlier_group','cluster']
        for row,record in zip(rows,records):
            for key,value in zip(keys,row):
                require(close(value,record[key]) if isinstance(record[key],(int,float,bool)) else value==record[key],f'Database mismatch {record["row_id"]}:{key}')
        result=connection.execute(benchmark['query'],benchmark['parameters']).fetchall()
        result_hash=hashlib.sha256(json.dumps(result,separators=(',',':')).encode()).hexdigest()
        require(result_hash==benchmark['before']['result_sha256']==benchmark['after']['result_sha256'],'SQL result hashes disagree.')
    validate_derivative(root,'dashboard',manifest)
    html=(root/'outputs/DASHBOARD_BAT_DONG_SAN.html').read_text(encoding='utf-8')
    require('__DATA_JSON__' not in html and manifest['run_id'] in html,'Dashboard provenance/data missing.')
    workbook_status={'status':'not_checked_optional_artifact'}
    if not core_only and (require_workbook or (root/'outputs/workbook_build.json').exists()):
        validate_derivative(root,'workbook',manifest)
        from openpyxl import load_workbook
        book=load_workbook(root/'outputs/DO_AN_BAT_DONG_SAN.xlsx',data_only=True,read_only=True)
        rows=list(book['Du lieu'].iter_rows(min_row=10,max_row=9+len(records),min_col=1,max_col=16,values_only=True))
        columns=['row_id','listing_id','province','district','property_type','transaction_status','area_m2','sale_price_vnd','price_per_m2_vnd','price_basis','area_source','quality_flags','is_outlier_global','is_outlier_group','iqr_group','cluster']
        require(len(rows)==len(records),'Workbook count mismatch.')
        for row,record in zip(rows,records):
            for key,value in zip(columns,row):
                require(close(value,record[key]) if isinstance(record[key],(int,float,bool)) else (value or '')==(record[key] or ''),f'Workbook mismatch {record["row_id"]}:{key}')
        expected_metrics=[len(records),len(records),analysis['overall']['median_price_vnd']/1e9,
                          analysis['overall']['mean_price_vnd']/1e9,analysis['overall']['median_area_m2'],
                          analysis['overall']['median_price_per_m2_vnd']/1e6]
        for row,expected in zip(range(7,13),expected_metrics):
            require(close(book['Dashboard'].cell(row,2).value,expected),f'Workbook KPI mismatch B{row}')
        require(sum(book['Dashboard'].cell(row,12).value for row in range(52,63))==len(records),'Workbook price bins do not cover all valid prices.')
        for sheet in book:
            for row in sheet.iter_rows():
                require(all(cell.data_type!='e' for cell in row),f'Cached spreadsheet error in {sheet.title}.')
        book.close();workbook_status={'status':'saved_values_and_provenance_verified','rows':len(rows)}
        evidence_path=root/'outputs/workbook_test_results.json'
        if evidence_path.exists():
            evidence=json.loads(evidence_path.read_text(encoding='utf-8'))
            require(evidence.get('core_run_id')==manifest['run_id'] and
                    evidence.get('workbook_sha256')==file_sha256(root/'outputs/DO_AN_BAT_DONG_SAN.xlsx') and
                    evidence.get('builder_sha256')==file_sha256(root/'tools/build_workbook.mjs'),
                    'Workbook recalculation evidence is stale.')
            workbook_status['recalculation_checks']={key:evidence[key] for key in ('status','scope','live_price_edit','live_area_edit','invalid_input_checks','chart_count')}
    python_tests=test_evidence(root,manifest['run_id'])
    dashboard_tests=dashboard_test_evidence(root,manifest['run_id'])
    require(python_tests['status']!='failed' and dashboard_tests['status']!='failed','Current test evidence contains failures.')
    report_status={'status':'not_verified_for_current_run','generated_this_run':False}
    report_meta=root/'outputs/report_build.json'
    if report_meta.exists():
        report=json.loads(report_meta.read_text(encoding='utf-8'))
        artifacts=report.get('artifacts',{})
        valid=bool(artifacts) and report.get('core_run_id')==manifest['run_id']
        for relative, expected in artifacts.items():
            candidate=(root/relative).resolve()
            valid=valid and candidate.is_relative_to(root.resolve()) and candidate.is_file()
            if valid:valid=file_sha256(candidate)==expected
        report_status={'status':'saved_report_hashes_verified' if valid else 'stale_or_missing_report',
                       'generated_this_run':bool(valid),'visual_review':report.get('visual_review','not_recorded')}
    return {'status':'passed','scope':'Saved core data, database and selected derivatives; read each test/visual/report status separately.',
            'core_run_id':manifest['run_id'],'raw_count':len(raw),'clean_count':len(records),
            'duplicates':analysis['counts']['duplicate_count'],'rejected':len(rejected),'source_sha256':raw_hash,
            'core_artifact_hashes_verified':len(manifest['artifacts']),'database_integrity':'ok','sql_results_identical':True,
            'dashboard':{'status':'saved_data_and_provenance_verified'},'workbook':workbook_status,
            'python_tests':python_tests,'dashboard_tests':dashboard_tests,
            'report':report_status,
            'visual_checks':'Not inferred by this command. See dated human/UI evidence separately.',
            'limitations':['Heuristic source-content checks cannot verify every source-entry error or transactions.',
                           'Human label validation is still necessary before strong conclusions.']}

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--core-only',action='store_true',help='Skip optional Excel explicitly.')
    parser.add_argument('--require-workbook',action='store_true',help='Require Excel regenerated from the current run.')
    args=parser.parse_args();destination=ROOT/'outputs/DELIVERY_CHECK.json'
    try:
        require(not(args.core_only and args.require_workbook),'Choose only one workbook verification mode.')
        result=verify(require_workbook=args.require_workbook,core_only=args.core_only)
    except Exception as exc:
        destination.parent.mkdir(parents=True,exist_ok=True)
        failure=json.dumps({'status':'failed','reason':str(exc)},ensure_ascii=False,indent=2)
        destination.write_text(failure,encoding='utf-8')
        (ROOT/'outputs/AUDIT_RESULTS.json').write_text(failure,encoding='utf-8')
        print('Verification FAILED: '+str(exc));return 1
    destination.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'outputs/AUDIT_RESULTS.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=True,indent=2));return 0

if __name__=='__main__': raise SystemExit(main())
