"""Build the six-chapter Vietnamese essay from authored prose and measured data.

Use the bundled document runtime. Core analysis and previous drafts are untouched.
"""
from __future__ import annotations
import csv
import hashlib
import json
import re
import sqlite3
from pathlib import Path
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
ASSETS=DOCS/'report_assets'
SOURCE=DOCS/'report_src/noi_dung.md'
TARGET=DOCS/'BAO_CAO_TIEU_LUAN_6_CHUONG.docx'
TITLE='Phân tích giá rao bất động sản Việt Nam'
BODY_WIDTH=17.09  # Letter page minus 2.5 cm left and 2 cm right margins.


def js(path):return json.loads((ROOT/path).read_text(encoding='utf-8'))
def csv_rows(path):
    with (ROOT/path).open(encoding='utf-8-sig',newline='') as stream:return list(csv.DictReader(stream))
def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def vn(n,d=0):return f'{float(n):,.{d}f}'.replace(',','_').replace('.',',').replace('_','.')
def million(n,d=4):return vn(float(n)/1e6,d)
def billion(n,d=4):return vn(float(n)/1e9,d)


def measured_tables(a):
    result={}
    def add(key,caption,headers,rows,widths,source):
        result[key]=dict(caption=caption,headers=headers,rows=rows,widths=widths,source=source)
    add('scope','Phạm vi ứng dụng và giới hạn suy luận',
        ['Thành phần','Phạm vi được chọn','Giới hạn'],[
        ['Đối tượng','Tin quảng cáo bất động sản trong tệp v2','Không đồng nhất tin với tài sản hoặc giao dịch'],
        ['Thời gian','Nguồn công bố 12/01/2021','Không suy ra giá thị trường 2026 hoặc chuỗi thời gian'],
        ['Địa phương','Nhãn địa chỉ lịch sử trong nguồn','Không suy tọa độ hay tự đổi địa giới hiện hành'],
        ['Người sử dụng','Người phân tích và người xem kết quả','Hai cách sử dụng, chưa có phân quyền tài khoản'],
        ['Đầu ra','Thống kê, hình, nhóm thăm dò, HTML/Excel/SQL','Chưa dùng định giá, môi giới hoặc khuyến nghị đầu tư']],
        [3.0,6.5,7.59],'Nguồn: thiết kế phạm vi nghiên cứu, đối chiếu metadata Mendeley v2 [1].')
    add('trace_design','Dữ liệu minh họa hai nhánh giữ và cách ly',
        ['Dòng nguồn','Dữ liệu quan sát','Đầu ra'],[
        ['1','100 m²; 2.380.000.000 VND','Giữ; 23.800.000 VND/m²'],
        ['638','143 m²; giá lõi 20 triệu; mô tả 20tr/m2','Cách ly do cơ sở giá chưa rõ; không tự sửa']],
        [2.0,8.0,7.09],'Nguồn: raw, listings_clean.csv, listings_rejected.csv và row_lineage.csv; cùng run_id.')
    fields=[
        ('row_id','Số nguyên','Vị trí nguồn từ 1; khóa truy nguyên'),
        ('listing_id','Chuỗi','Tiền tố hash nội dung; không phải ID tài sản'),
        ('province','Chuỗi','Nhãn tỉnh xác định hoặc Chưa xác định'),
        ('district','Chuỗi','Thành phần địa chỉ theo mẫu; không suy từ đường'),
        ('property_type','Chuỗi','land / house / apartment / other_unknown'),
        ('transaction_status','Chuỗi','sale / sale_assumed / sale_with_rental_income'),
        ('area_m2','Số, m²','Diện tích được đọc; không chứng thực đất hay sàn'),
        ('sale_price_vnd','Số, VND','Giá rao tổng chuẩn hóa theo cơ sở giá'),
        ('price_per_m2_vnd','Số, VND/m²','Giá tổng chia diện tích từng tin'),
        ('price_basis','Chuỗi','total hoặc per_m2; giữ cơ sở quy đổi'),
        ('area_source','Chuỗi','Trường diện tích hoặc phần ghi rõ trong chuỗi giá'),
        ('source_review_evidence','JSON dạng chuỗi','Cụm số/chữ của quy tắc kiểm, không phải nhãn chuẩn'),
        ('quality_flags','Chuỗi','Các cờ chất lượng, phân cách bằng dấu chấm phẩy'),
        ('is_outlier_global','Boolean','Cờ IQR toàn mẫu'),
        ('is_outlier_group','Boolean','Cờ IQR theo nhóm có dự phòng'),
        ('iqr_group','Chuỗi','Nhóm tham chiếu dùng tính ngưỡng'),
        ('cluster','Số nguyên','1 hoặc 2; −1 cho tin ngoài tập fit'),
        ('area_band','Chuỗi','Khoảng diện tích có biên xác định')]
    add('dictionary','Từ điển 18 biến của dữ liệu sạch',['Biến','Kiểu và đơn vị','Ý nghĩa'],fields,[5.1,3.0,8.99],
        'Nguồn: schema listings_clean.csv và src/pipeline.py. CSV giữ 6.936 tin.')
    conn=sqlite3.connect((ROOT/'outputs/real_estate.sqlite').as_uri()+'?mode=ro',uri=True)
    row=conn.execute('SELECT l.row_id,l.province_id,p.name,l.type_id,t.name,l.area_m2,l.sale_price_vnd,l.price_per_m2_vnd FROM listings l JOIN provinces p USING(province_id) JOIN property_types t USING(type_id) WHERE row_id=1').fetchone()
    add('db_example','Bản ghi minh họa liên kết ba bảng', ['Đối tượng','Khóa và tên','Dữ liệu minh họa'],[
        ['provinces',f'province_id={row[1]}; {row[2]}','48 hàng, gồm nhóm chưa rõ'],
        ['property_types',f'type_id={row[3]}; {row[4]}','4 hàng nhãn suy luận'],
        ['listings',f'row_id={row[0]}; tham chiếu hai khóa trên',f'{vn(row[5])} m²; {billion(row[6],2)} tỷ; {million(row[7],1)} triệu/m²']],
        [3.1,7.0,6.99],'Nguồn: truy vấn chỉ đọc outputs/real_estate.sqlite; lược đồ sql/01_schema.sql.')
    add('modules','Liên kết mô-đun với đầu vào và đầu ra',['Mô-đun','Đầu vào','Đầu ra'],[
        ['download_data.py','URL pin và manifest nguồn','JSON nguyên bản, metadata'],
        ['pipeline.py / source_checks.py','Năm trường lõi của raw','CSV/JSON sạch, rejected, lineage, bảng/hình/DB'],
        ['artifact_contract.py','Nguồn, mã, manifest và sản phẩm','Xác nhận phiên bản hoặc lỗi'],
        ['build_dashboard.py / template','CSV sạch và manifest','HTML có dữ liệu nhúng, metadata dựng'],
        ['build_workbook.mjs','Dữ liệu sạch và analysis','Excel ba trang/bốn hình, metadata/kiểm tính lại'],
        ['demo_sql.py','SQLite và bốn SQL','Kết quả truy vấn chỉ đọc'],
        ['verify_delivery.py','Các sản phẩm cùng phiên bản','Kết quả đối soát hoặc lỗi'],
        ['run_project.py','Cấu hình chạy và nguồn','Điều phối toàn bộ phần được chọn']],
        [5.3,5.3,6.49],'Nguồn: mã nguồn trong src, scripts và tools tại phiên bản phân tích [17].')
    add('operation','Lệnh vận hành và phạm vi tạo kết quả',['Lệnh từ thư mục repo','Phạm vi'],[
        [r'python scripts/run_project.py','Tải/kiểm nguồn, chạy lõi, dựng HTML, test/SQL/verifier; Excel không tự dựng mặc định'],
        [r'python scripts/run_project.py --skip-download','Dùng raw có sẵn, vẫn kiểm checksum'],
        [r'python scripts/run_project.py --workbook','Thêm bước dựng Excel; cần Node và Artifact Tool'],
        [r'python tools/verify_delivery.py --require-workbook','Đối soát sản phẩm lưu gồm Excel, không chạy lại phân tích'],
        [r'python scripts/demo_sql.py','Bốn truy vấn đọc và kiểm toàn vẹn'],
        [r'MO_DASHBOARD.cmd','Mở HTML có sẵn trong trình duyệt']],
        [8.2,8.89],'Nguồn: scripts/run_project.py và README. Dùng Python trong môi trường đã cài requirements.')
    c=a['counts']
    add('counts','Đối soát quy mô dữ liệu trước và sau xử lý',['Trạng thái','Số dòng','Ý nghĩa'],[
        ['Nguồn',vn(c['raw_count']),'Mảng JSON công bố v2'],
        ['Trùng',vn(c['duplicate_count']),'Trùng theo năm trường lõi đã chuẩn hóa'],
        ['Cách ly',vn(c['rejected_count']),'Chưa đáp ứng chính sách, không đồng nghĩa tin giả'],
        ['Giữ',vn(c['clean_count']),'Tập EDA mặc định'],
        ['Ngoại lệ IQR nhóm',vn(c['group_iqr_outlier_count']),'Vẫn thuộc tập giữ mặc định'],
        ['Fit K-means',vn(a['clustering']['fit_count']),'Tập giữ sau loại cờ IQR nhóm']],
        [4.1,2.5,10.49],'Nguồn: outputs/analysis.json, counts và clustering; đơn vị dòng/tin.')
    desc=csv_rows('outputs/tables/descriptive_statistics.csv')
    add('descriptive','Thống kê minh họa phân bố giá và diện tích',
        ['Chỉ tiêu','Diện tích m²','Giá tổng tỷ VND','Đơn giá triệu VND/m²'],[
        [label, vn(desc[0][key],2),billion(desc[1][key],4),million(desc[2][key],4)] for key,label in [('count','Số tin'),('mean','Trung bình'),('25%','Q1'),('50%','Trung vị'),('75%','Q3'),('max','Lớn nhất')]],
        [3.5,3.6,4.7,5.29],'Nguồn: descriptive_statistics.csv; hàng số tin ở ba cột đều 6.936, không áp dụng quy đổi đơn vị.')
    result['descriptive']['rows'][0]=['Số tin','6.936','6.936','6.936']
    clean=csv_rows('data/processed/listings_clean.csv')
    examples=[r for r in clean if int(r['row_id']) in {1,2,3,4,5}]
    add('scatter_examples','Một số tin minh họa cho biểu đồ diện tích và giá',['Dòng','Loại suy luận','Diện tích m²','Giá tổng tỷ VND'],[
        [r['row_id'],r['property_type'],vn(r['area_m2'],2),billion(r['sale_price_vnd'],4)] for r in examples],
        [1.5,5.0,4.2,6.39],'Nguồn: listings_clean.csv; các dòng minh họa thuộc hình scatter toàn bộ 6.936 tin.')
    spearman=a['spearman']
    add('correlations','Hệ số Spearman giữa các biến phân tích',['Cặp biến','Hệ số','Phạm vi diễn giải'],[
        ['Diện tích – tổng giá',vn(spearman['area_m2']['sale_price_vnd'],6),'Mối liên hệ đơn điệu tổng thể trong mẫu'],
        ['Diện tích – đơn giá',vn(spearman['area_m2']['price_per_m2_vnd'],6),'Có liên hệ toán học qua mẫu số diện tích'],
        ['Tổng giá – đơn giá',vn(spearman['sale_price_vnd']['price_per_m2_vnd'],6),'Có biến giá chung trong tử số']],
        [5.7,2.5,8.89],'Nguồn: analysis.json, spearman; n=6.936. Không diễn giải nhân quả.')
    add('province_top','Dữ liệu mười địa phương có nhiều tin nhất trong hình',
        ['Địa phương','n','Trung vị triệu VND/m²','Bootstrap 95% triệu VND/m²'],[
        [r['province'],vn(r['n']),million(r['median_price_per_m2_vnd'],2),f"{million(r['bootstrap_low_vnd'],2)} – {million(r['bootstrap_high_vnd'],2)}"] for r in a['province_summary'][:10]],
        [5.0,1.7,5.0,5.39],'Nguồn: analysis.json, province_summary. Chọn theo số tin; khoảng bootstrap có điều kiện, 500 lần.')
    labels={'house':'Nhà','land':'Đất','apartment':'Căn hộ','other_unknown':'Chưa xác định'}
    add('types','Thống kê theo loại hình suy luận từ tiêu đề',
        ['Loại suy luận','n','Diện tích trung vị m²','Đơn giá trung vị triệu/m²'],[
        [labels[r['property_type']],vn(r['n']),vn(r['median_area_m2'],2),million(r['median_price_per_m2_vnd'],4)] for r in a['type_summary']],
        [4.1,2.0,5.0,5.99],'Nguồn: analysis.json, type_summary; nhãn chưa có kiểm định bằng bộ nhãn chuẩn.')
    scenarios=['Toàn bộ tin giữ','Bỏ IQR toàn cục','Bỏ IQR nhóm','Bỏ dấu chấm nghìn','Tiêu đề bán rõ']
    add('sensitivity','Độ nhạy của trung vị đơn giá theo lựa chọn xử lý',
        ['Phương án','n','Trung vị triệu/m²','Thay đổi %'],[
        [name,vn(r['n']),million(r['median_price_per_m2_vnd'],4),vn(r['unit_price_median_change_pct'],2)] for name,r in zip(scenarios,a['sensitivity'])],
        [6.0,2.0,5.2,3.89],'Nguồn: analysis.json, sensitivity; so với toàn bộ 6.936 tin giữ. Các phương án khác nhau về thành phần mẫu.')
    cl=a['clustering']
    add('cluster_trials','Khảo sát số cụm trên cùng tập fit', ['k','Silhouette mẫu','Inertia','Cụm nhỏ nhất'],[
        [str(r['k']),vn(r['silhouette_sample'],4),vn(r['inertia'],4),vn(r['min_cluster_size'])] for r in cl['trials']],
        [1.5,5.0,5.5,5.09],'Nguồn: cluster_trials.csv và analysis.json; fit 6.676, silhouette lấy tối đa 1.000 tin, seed 20261015.')
    add('cluster_summary','Đặc điểm hai cụm ở đơn vị gốc', ['Cụm','n','Diện tích trung vị m²','Đơn giá trung vị triệu/m²'],[
        [str(r['cluster']),vn(r['n']),vn(r['median_area_m2']),million(r['median_price_per_m2_vnd'],4)] for r in cl['summary']],
        [1.5,2.2,6.3,7.09],'Nguồn: cluster_summary.csv và analysis.json; nhãn cụm phục vụ thăm dò, chưa xác nhận phân khúc nghiệp vụ.')
    benchmark=a['sql_benchmark']
    rows=conn.execute(benchmark['query'],benchmark['parameters']).fetchall()
    assert hashlib.sha256(json.dumps(rows,separators=(',',':')).encode()).hexdigest()==benchmark['after']['result_sha256']
    add('sql_results','Năm dòng đầu của truy vấn khoảng đơn giá', ['row_id','Diện tích m²','Giá tỷ VND','Đơn giá triệu/m²'],[
        [str(r[0]),vn(r[1],2),billion(r[2],4),million(r[3],4)] for r in rows[:5]],
        [2.1,4.3,4.8,5.89],f"Nguồn: SELECT thật từ SQLite; TP.HCM, house, đơn giá từ {million(benchmark['parameters'][2],4)} đến {million(benchmark['parameters'][3],4)} triệu/m²; LIMIT 100.")
    conn.close()
    add('sql_benchmark','Thời gian truy vấn trước và sau chỉ mục', ['Giai đoạn','Trung vị ms','P10 ms','P90 ms'],[
        [label,vn(benchmark[key]['median_ms'],4),vn(benchmark[key]['p10_ms'],4),vn(benchmark[key]['p90_ms'],4)] for key,label in [('before','Trước index'),('after','Sau index')]],
        [4.5,4.2,4.2,4.19],'Nguồn: sql_benchmark.json; cùng truy vấn/tham số, mỗi pha 10 warmup + 51 lần đo; thứ tự tuần tự warm-cache.')
    add('dashboard_states','Kết quả kiểm logic các trạng thái dashboard', ['Trạng thái','Số tin','Phạm vi kiểm'],[
        ['Mặc định','6.936','JavaScript thật trên DOM tối giản'],
        ['TP.HCM','3.277','Như trên'],['TP.HCM + căn hộ','239','Như trên'],
        ['Bỏ IQR nhóm','6.676','Như trên'],['Chỉ tiêu đề bán rõ','3.030','Như trên'],
        ['Điều kiện không có dữ liệu','0','Ca kiểm lập trình; chưa khẳng định thao tác được từ mọi UI'],
        ['Tìm Hà Nội trong bảng','1.014','Bảng/CSV; KPI toàn trang giữ 6.936 khi lọc trang mặc định']],
        [6.0,2.0,9.09],'Nguồn: dashboard_test_results.json và AUDIT_RESULTS.json; chưa thay thế kiểm bố cục trình duyệt bản mới.')
    add('verification','Phạm vi và kết quả các lớp kiểm chứng', ['Lớp kiểm','Bằng chứng','Kết quả và giới hạn'],[
        ['Mã Python','29 test; 0 lỗi; 0 bỏ qua','Đạt theo bộ test hiện có'],
        ['Dữ liệu/phiên bản','Verifier so từng bản ghi, hash, lineage','CSV/JSON/DB/workbook cùng dữ liệu'],
        ['Trung vị độc lập','csv/statistics và SQLite ORDER BY','Ba trung vị khớp analysis.json'],
        ['SQL','Bốn truy vấn; integrity/FK; hash kết quả','Có DB thật; index không tăng tốc ở lần ghi nhận'],
        ['HTML','JS thật với DOM mô phỏng','Logic đạt; bố cục trình duyệt mới chưa xác nhận'],
        ['Excel','6.936 dòng; 4 biểu đồ; thử tính lại engine','Artifact Tool đạt; Excel thật chưa xác nhận'],
        ['Tái chạy','Bản sao nguồn cùng Python đã cài','6.936 dòng giống; chưa cài mới trên máy khác'],
        ['Nội dung nguồn','60 tin + 7 tin trọng điểm do trợ lý đọc','Phát hiện giới hạn; không phải gold labels']],
        [3.3,6.2,7.59],'Nguồn: AUDIT_RESULTS.json, ASSISTANT_CHECKS.json và các log kiểm/tái lập cùng run_id.')
    add('conclusion','Mức đáp ứng mục tiêu theo bằng chứng hiện có', ['Khía cạnh','Mức đã đạt','Giới hạn cần giữ'],[
        ['Tin cậy','Đúng bản nguồn, tính toán/phiên bản đối soát','Chưa xác minh mọi tin hoặc độ chính xác nhãn'],
        ['Vận hành','Lõi, SQL, HTML và workbook đã được tạo/kiểm','Thiếu xác nhận UI mới, Excel thật và cài máy mới'],
        ['Tích hợp','Dữ liệu chung, metadata, một lệnh điều phối','Tích hợp tệp cục bộ; workbook phụ thuộc runtime riêng'],
        ['Ứng dụng','Khám phá mẫu lịch sử, học tập, truy nguyên','Chưa dùng định giá hiện tại, pháp lý hoặc đầu tư']],
        [3.0,6.6,7.49],'Nguồn: tổng hợp kết quả và giới hạn ở Chương 5; phạm vi mục tiêu Chương 1.')
    return result


FIGURES={
    'architecture':dict(caption='Sơ đồ luồng xử lý và các sản phẩm dữ liệu',path='docs/report_assets/architecture.png',note='Dữ liệu minh họa: 9.004 nguồn → 1.884 trùng + 184 cách ly + 6.936 giữ; dòng 1 tạo 23,8 triệu/m², dòng 638 cách ly.'),
    'erd':dict(caption='Sơ đồ quan hệ ba bảng trong SQLite',path='docs/report_assets/erd.png',note='Dữ liệu minh họa: provinces 48 hàng, property_types 4 hàng, listings 6.936 hàng; ví dụ liên kết cụ thể tại Bảng [[table:db_example]].'),
    '01_price_distribution':dict(caption='Phân bố tổng giá rao ở thang gốc và thang log',note='n=6.936; tổng giá theo tỷ VND, log10(giá/tỷ VND); lần lượt 50 và 35 khoảng. Số liệu minh họa tại Bảng [[table:descriptive]].'),
    '02_area_price_scatter':dict(caption='Mối liên hệ giữa diện tích và tổng giá rao',note='n=6.936, hai trục log; diện tích m² và tổng giá tỷ VND, màu theo loại suy luận. Ví dụ tại Bảng [[table:scatter_examples]].'),
    '08_spearman_correlation':dict(caption='Ma trận tương quan Spearman của ba biến phân tích',note='n=6.936; hệ số không có đơn vị, tính từ thứ hạng. Bảng [[table:correlations]] nêu các cặp biến, không diễn giải nhân quả.'),
    '03_province_medians':dict(caption='Trung vị đơn giá của mười địa phương có nhiều tin nhất',note='Chọn theo n, chỉ nhóm n≥30; đơn vị triệu VND/m²; bootstrap 95% có điều kiện, 500 lần. Bảng [[table:province_top]] kèm toàn bộ số liệu hình.'),
    '06_cleaning_sensitivity':dict(caption='Đối soát làm sạch và độ nhạy theo lựa chọn xử lý',note='Phần trái đơn vị dòng/tin; phần phải trung vị triệu VND/m². Số liệu tương ứng ở Bảng [[table:counts]] và Bảng [[table:sensitivity]].'),
    '07_exploratory_clusters':dict(caption='Phân nhóm K-means và khảo sát số cụm',note='Fit 6.676 tin; diện tích m², đơn giá triệu VND/m² trên trục log; silhouette lấy mẫu tối đa 1.000 tin. Dữ liệu tại Bảng [[table:cluster_trials]] và Bảng [[table:cluster_summary]].'),
    'workbook_kpi':dict(caption='Vùng chỉ tiêu trong bản render trang Dashboard của Excel',path='docs/report_assets/workbook_kpi.png',note='Ảnh render bằng Artifact Tool, không phải thao tác Microsoft Excel. Dữ liệu: 6.936 tin; trung vị 2,25 tỷ, trung bình hiển thị 3,70 tỷ; diện tích 70 m²; đơn giá 32,32 triệu/m². Đối chiếu Bảng [[table:descriptive]].'),
}


def diagram_assets():
    ASSETS.mkdir(parents=True,exist_ok=True)
    preview=ROOT/'tmp/qa_workbook/dashboard.png'
    if preview.exists():
        assert sha(preview)=='2fbaa4de77f68505c2d0aa5586535e3f4efb66b2fb38cdf99148db20b07b6de1'
        image=Image.open(preview)
        image.crop((60,250,636,585)).save(ASSETS/'workbook_kpi.png')
    else:
        assert (ASSETS/'workbook_kpi.png').is_file(), 'Committed workbook preview is required to reproduce this report.'
    font_dir=Path('C:/Windows/Fonts')
    regular=font_dir/'arial.ttf';bold=font_dir/'arialbd.ttf'
    def font(size,heavy=False):return ImageFont.truetype(str(bold if heavy else regular),size)
    def box(draw,xy,title,detail):
        draw.rounded_rectangle(xy,14,fill='#f2f5f7',outline='#607789',width=3)
        x0,y0,x1,y1=xy
        draw.text(((x0+x1)/2,y0+28),title,font=font(31,True),anchor='mt',fill='black')
        for i,line in enumerate(detail.splitlines()):
            draw.text(((x0+x1)/2,y0+80+i*43),line,font=font(27),anchor='mt',fill='#172c40')
    def arrow(draw,start,end):
        import math
        draw.line([start,end],fill='#426278',width=5)
        theta=math.atan2(end[1]-start[1],end[0]-start[0]);length=20
        p1=(end[0]-length*math.cos(theta-.45),end[1]-length*math.sin(theta-.45))
        p2=(end[0]-length*math.cos(theta+.45),end[1]-length*math.sin(theta+.45))
        draw.polygon([end,p1,p2],fill='#426278')
    im=Image.new('RGB',(1800,1030),'white');d=ImageDraw.Draw(im)
    box(d,(35,40,575,225),'Nguồn pin SHA-256','9.004 bản ghi JSON\nMetadata và manifest nguồn')
    box(d,(650,40,1190,225),'Chuẩn bị dữ liệu','Parse · trùng · nhất quán\nGiữ raw và truy nguyên')
    box(d,(1265,40,1765,225),'Nhánh ngoài EDA','1.884 dòng trùng\n184 tin cách ly')
    arrow(d,(575,130),(650,130));arrow(d,(1190,130),(1265,130))
    box(d,(650,315,1190,500),'Dữ liệu đủ điều kiện','6.936 tin giữ\nCSV · JSON · lineage')
    arrow(d,(920,225),(920,315))
    box(d,(35,590,575,800),'Thống kê và hình','IQR · trung vị · độ nhạy\nK-means fit 6.676 tin')
    box(d,(650,590,1190,800),'Cơ sở dữ liệu','SQLite ba bảng\nSQL và benchmark thật')
    box(d,(1265,590,1765,800),'Sản phẩm xem dữ liệu','HTML có dữ liệu nhúng\nExcel ba trang, bốn hình')
    arrow(d,(780,500),(400,590));arrow(d,(920,500),(920,590));arrow(d,(1060,500),(1500,590))
    d.rounded_rectangle((240,880,1560,990),12,fill='#e7eef1',outline='#607789',width=2)
    d.text((900,910),'Manifest và verifier kiểm nguồn, mã, dữ liệu và sản phẩm cùng phiên bản',font=font(29),anchor='mt',fill='black')
    im.save(ASSETS/'architecture.png')
    im=Image.new('RGB',(1800,1210),'white');d=ImageDraw.Draw(im)
    def entity(xy,title,fields):
        x0,y0,x1,y1=xy
        d.rectangle(xy,fill='white',outline='#607789',width=3)
        d.rectangle((x0,y0,x1,y0+75),fill='#e1e9ef',outline='#607789',width=2)
        d.text((x0+25,y0+20),title,font=font(34,True),fill='black')
        for i,line in enumerate(fields):d.text((x0+25,y0+103+i*51),line,font=font(29),fill='black')
    entity((35,100,560,390),'provinces · 48 hàng',['PK  province_id','UQ  name'])
    entity((35,610,560,900),'property_types · 4 hàng',['PK  type_id','UQ  name'])
    entity((850,45,1760,1050),'listings · 6.936 hàng',[
        'PK  row_id','UQ  listing_id','FK  province_id','FK  type_id','district','transaction_status',
        'area_m2 > 0','sale_price_vnd ≥ 10.000.000','price_per_m2_vnd > 0','is_outlier_group: 0 / 1','cluster: -1 / 1 / 2 trong bản này'])
    d.line([(560,245),(700,245),(700,275),(850,275)],fill='#426278',width=4)
    d.text((590,192),'1',font=font(30),fill='black');d.text((795,223),'N',font=font(30),fill='black')
    d.line([(560,755),(720,755),(720,325),(850,325)],fill='#426278',width=4)
    d.text((590,701),'1',font=font(30),fill='black');d.text((795,330),'N',font=font(30),fill='black')
    d.text((55,1090),'PK khóa chính   FK khóa ngoại   UQ duy nhất',font=font(26),fill='#334a5b')
    d.text((55,1140),'Cluster liệt kê giá trị bản chạy; lược đồ không có CHECK enum cho cột này.',font=font(26),fill='#334a5b')
    im.save(ASSETS/'erd.png')


# Native Office Math builders. No LaTeX strings are inserted into the report.
def mr(text):
    node=OxmlElement('m:r');pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);node.append(pr)
    t=OxmlElement('m:t');t.text=text;node.append(t);return node
def wrapper(tag,items):
    node=OxmlElement(tag)
    for item in items:node.append(item)
    return node
def frac(a,b):return wrapper('m:f',[wrapper('m:num',a),wrapper('m:den',b)])
def sub(base,lower):return wrapper('m:sSub',[wrapper('m:e',[mr(base)]),wrapper('m:sub',[mr(lower)])])
def sup(base,upper):return wrapper('m:sSup',[wrapper('m:e',[mr(base)]),wrapper('m:sup',[mr(upper)])])
def subsup(base,lower,upper):return wrapper('m:sSubSup',[wrapper('m:e',[mr(base)]),wrapper('m:sub',[mr(lower)]),wrapper('m:sup',[mr(upper)])])
def summation(lower,upper,items):
    node=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');char=OxmlElement('m:chr');char.set(qn('m:val'),'∑');pr.append(char)
    lim=OxmlElement('m:limLoc');lim.set(qn('m:val'),'subSup');pr.append(lim);node.append(pr)
    node.append(wrapper('m:sub',[mr(lower)]));node.append(wrapper('m:sup',[mr(upper)]));node.append(wrapper('m:e',items));return node
def equation_nodes(key):
    if key=='unit_price':return [sub('u','i'),mr(' = '),frac([sub('P','i')],[sub('A','i')]),mr(' ,  '),sub('A','i'),mr(' > 0')]
    if key=='iqr':return [mr('IQR = '),sub('Q','3'),mr(' − '),sub('Q','1'),mr(' ;  L = '),sub('Q','1'),mr(' − 1,5 IQR ;  U = '),sub('Q','3'),mr(' + 1,5 IQR')]
    if key=='spearman':return [sub('ρ','s'),mr(' = corr(R(A), R(P))')]
    if key=='bootstrap':return [mr('CI = ['),subsup('q','0,025','*'),mr(' , '),subsup('q','0,975','*'),mr('] ;  B = 500')]
    if key=='zscore':return [sub('z','ij'),mr(' = '),frac([mr('ln('),sub('x','ij'),mr(') − '),sub('μ','j')],[sub('σ','j')]),mr(' ;  j ∈ {diện tích, đơn giá}')]
    if key=='kmeans':return [mr('J = '),summation('j=1','k',[summation('i∈Cⱼ','',[mr('‖'),sub('z','i'),mr(' − '),sub('c','j'),sup('‖','2')])])]
    if key=='centroid':return [sub('c','j'),mr(' = '),frac([mr('1')],[mr('|'),sub('C','j'),mr('|')]),summation('i∈Cⱼ','',[sub('z','i')])]
    if key=='silhouette':return [sub('s','i'),mr(' = '),frac([sub('b','i'),mr(' − '),sub('a','i')],[mr('max('),sub('a','i'),mr(', '),sub('b','i'),mr(')')])]
    raise KeyError(key)


def main():
    a=js('outputs/analysis.json');assert a['run_id']=='20261016T015620Z-f2e661a6'
    manifest=js('outputs/run_manifest.json')
    for relative,digest in manifest['artifacts'].items():assert sha(ROOT/relative)==digest,relative
    for relative,digest in manifest['source_files'].items():assert sha(ROOT/relative)==digest,relative
    for name in ['dashboard','workbook']:
        metadata=js(f'outputs/{name}_build.json');assert metadata['core_run_id']==a['run_id'];assert sha(ROOT/metadata['path'])==metadata['sha256']
    diagram_assets();tables=measured_tables(a)
    source=SOURCE.read_text(encoding='utf-8');chunks=source.strip().split('\n\n')
    indices={'table':{},'figure':{},'equation':{}}
    for chunk in chunks:
        if chunk.startswith('@'):
            kind,key=chunk[1:].strip().split(' ',1)
            if kind in indices:indices[kind][key]=len(indices[kind])+1
    figure_count=len(indices['figure']);table_count=len(indices['table'])
    def resolve(text):
        def replacement(match):return str(indices[match[1]][match[2]])
        return re.sub(r'\[\[(table|figure|equation):([^\]]+)\]\]',replacement,text)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.0)
    sec.header_distance=Cm(.8);sec.footer_distance=Cm(.9);sec.different_first_page_header_footer=True
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
        for k in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+k),'Times New Roman')
    for st in doc.styles:
        pp=st.element.find(qn('w:pPr'))
        if pp is not None:
            for item in list(pp):
                if item.tag==qn('w:pBdr'):pp.remove(item)
        rp=st.element.find(qn('w:rPr'))
        if rp is not None:
            for fonts in rp.findall(qn('w:rFonts')):
                for key in list(fonts.attrib):
                    if 'Theme' in key:del fonts.attrib[key]
    normal=doc.styles['Normal'];normal.font.size=Pt(12)
    pf=normal.paragraph_format;pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.line_spacing=1.15;pf.space_after=Pt(6);pf.first_line_indent=Cm(.7);pf.widow_control=True
    for name,size in [('Title',22),('Subtitle',14),('Heading 1',16),('Heading 2',13),('Caption',11)]:doc.styles[name].font.size=Pt(size)
    doc.styles['Heading 1'].font.bold=True
    hp=doc.styles['Heading 1'].paragraph_format;hp.page_break_before=True;hp.keep_with_next=True;hp.space_before=Pt(0);hp.space_after=Pt(15);hp.first_line_indent=Cm(0)
    doc.styles['Caption'].font.italic=False
    settings=doc.settings.element
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
    mathpr=OxmlElement('m:mathPr');mf=OxmlElement('m:mathFont');mf.set(qn('m:val'),'Cambria Math');mathpr.append(mf);settings.append(mathpr)
    header=sec.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.CENTER;header.paragraph_format.first_line_indent=Cm(0)
    run=header.add_run('NHẬP MÔN KHOA HỌC DỮ LIỆU  |  PHÂN TÍCH GIÁ RAO BẤT ĐỘNG SẢN');run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Cm(0)
    footer.add_run('Trang ').font.size=Pt(10)
    def field(paragraph,instruction,result='0'):
        start=OxmlElement('w:fldChar');start.set(qn('w:fldCharType'),'begin');paragraph.add_run()._r.append(start)
        instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' '+instruction+' ';paragraph.add_run()._r.append(instr)
        sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');paragraph.add_run()._r.append(sep)
        paragraph.add_run(str(result));end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');paragraph.add_run()._r.append(end)
    field(footer,'PAGE','2')
    bookmarks=[]
    def bookmark(p,name):
        i=len(bookmarks)+1;bookmarks.append(name)
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(i));start.set(qn('w:name'),name)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(i));p._p.insert(0,start);p._p.append(end)
    def p(text,style=None,center=False,small=False):
        para=doc.add_paragraph(resolve(text),style)
        if center:para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.first_line_indent=Cm(0)
        if small:
            para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.line_spacing=1.0;para.paragraph_format.space_after=Pt(7)
            for run in para.runs:run.font.size=Pt(10.5)
        return para
    def front_title(text,newpage=True):
        if newpage:doc.add_page_break()
        para=p(text,center=True);para.paragraph_format.space_after=Pt(18)
        for run in para.runs:run.bold=True;run.font.size=Pt(16)
        return para
    # Cover contains only information supplied by the user; no invented campus/class.
    for text,size,bold,after in [
        ('HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG',15,True,35),
        ('TIỂU LUẬN HỌC PHẦN',15,True,12),('NHẬP MÔN KHOA HỌC DỮ LIỆU',17,True,32)]:
        para=p(text,center=True);para.paragraph_format.space_after=Pt(after)
        for r in para.runs:r.font.size=Pt(size);r.bold=bold
    para=p(TITLE,'Title',center=True);para.paragraph_format.space_after=Pt(38)
    for r in para.runs:
        r.font.name='Times New Roman';r.font.size=Pt(23);r.bold=True;r.font.color.rgb=RGBColor(0,0,0)
        fonts=r._r.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if 'Theme' in key:del fonts.attrib[key]
        fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman')
    p('Giảng viên hướng dẫn  Nguyễn Ngọc Duy',center=True)
    p('Nhóm thực hiện',center=True)
    for name,student in [('Võ Văn Đạt','N23DCCN078'),('Nguyễn Thị Yến Nhi','N23DCCN112'),('Hứa Như Quỳnh','N23DCCN118')]:p(f'{name}  —  {student}',center=True)
    para=p('Năm 2026',center=True);para.paragraph_format.space_before=Pt(45)
    front_title('Tóm tắt')
    p('Tiểu luận xây dựng một quy trình có thể chạy lại để khám phá và so sánh giá rao trong mẫu tin bất động sản lịch sử từ Vietnam Real Estate Data v2, công bố trên Mendeley Data năm 2021. Từ 9.004 bản ghi, quy trình gộp 1.884 dòng trùng theo năm trường lõi, cách ly 184 dòng chưa đáp ứng chính sách và giữ 6.936 tin. Bản nguồn được kiểm checksum; dữ liệu sạch đi kèm cờ chất lượng và nhật ký truy nguyên để giải thích các quyết định xử lý.')
    p('Kết quả cho thấy trung vị tổng giá 2,25 tỷ VND, diện tích 70 m² và đơn giá khoảng 32,3205 triệu VND/m². Mẫu tập trung nhiều ở TP.HCM và Hà Nội, còn cách ghi số và lựa chọn tập lọc có thể làm thay đổi chỉ tiêu. K-means được lập trình bằng NumPy trên hai biến log đã chuẩn hóa, tạo hai nhóm thăm dò từ 6.676 tin. Cơ sở dữ liệu SQLite ba bảng cho phép truy vấn thật và khảo sát index; lần đo ghi nhận thay đổi kế hoạch nhưng chưa cải thiện độ trễ trung vị.')
    p('Sản phẩm gồm bảng/hình, dashboard HTML ngoại tuyến và workbook Excel. Kiểm thử Python, logic JavaScript và đối soát đa định dạng hỗ trợ tính nhất quán kỹ thuật. Nghiên cứu vẫn giới hạn bởi nội dung tin chưa xác minh, nhãn suy luận, mẫu thuận tiện và thiếu ngày từng tin. Kết quả phù hợp khám phá dữ liệu và học tập trong phạm vi lịch sử đã chọn, chưa dùng để định giá thị trường hiện tại hoặc khuyến nghị đầu tư.')
    p('Từ khóa: giá rao bất động sản; chuẩn bị dữ liệu; trực quan hóa; truy nguyên; K-means; SQLite.',small=True)
    p('Mốc thực nghiệm: 20261016T015620Z-f2e661a6. Nguồn dữ liệu công bố 12/01/2021; thời điểm xây dựng tiểu luận không phải thời điểm giá thị trường.',small=True)
    chapters=[chunk[2:] for chunk in chunks if chunk.startswith('# ')]
    front_title('Mục lục')
    toc_entries=[(title,f'chapter_{i}') for i,title in enumerate(chapters,1)]+[('Tài liệu tham khảo','references'),('Phụ lục tái lập và nguồn minh chứng','appendix')]
    def pageref_line(label,target):
        para=p(label);para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.line_spacing=1.0;para.paragraph_format.space_after=Pt(8)
        tabs=para.paragraph_format.tab_stops
        from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
        tabs.add_tab_stop(Cm(BODY_WIDTH),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
        para.add_run('\t');field(para,'PAGEREF '+target+' \\h')
        return para
    for label,target in toc_entries:pageref_line(label,target)
    front_title('Danh mục hình và bảng')
    for key,number in indices['figure'].items():pageref_line(f"Hình {number}: {FIGURES[key]['caption']}",f'figure_{number}')
    for key,number in indices['table'].items():pageref_line(f"Bảng {number}: {tables[key]['caption']}",f'table_{number}')
    markdown=['# '+TITLE,'','Học viện Công nghệ Bưu chính Viễn thông','Học phần: Nhập môn khoa học dữ liệu','Giảng viên: Nguyễn Ngọc Duy','']
    chapter_no=0
    for chunk in chunks:
        if chunk.startswith('# '):
            chapter_no+=1;para=p(chunk[2:],'Heading 1');bookmark(para,f'chapter_{chapter_no}');markdown.extend([chunk,'']);continue
        if not chunk.startswith('@'):
            p(chunk);markdown.extend([resolve(chunk),'']);continue
        kind,key=chunk[1:].strip().split(' ',1)
        if kind=='table':
            t=tables[key];number=indices['table'][key];caption=f"Bảng {number}: {t['caption']}"
            para=p(caption,'Caption',center=True);para.paragraph_format.keep_with_next=True;bookmark(para,f'table_{number}')
            table=doc.add_table(rows=1,cols=len(t['headers']));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            total=sum(t['widths']);widths=[x/total*BODY_WIDTH for x in t['widths']]
            for col,w in zip(table.columns,widths):col.width=Cm(w)
            for rowvalues in [t['headers']]+t['rows']:
                row=table.rows[0] if rowvalues is t['headers'] else table.add_row()
                for i,(cell,text) in enumerate(zip(row.cells,rowvalues)):
                    cell.width=Cm(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text=resolve(str(text));cp=cell.paragraphs[0];cp.paragraph_format.first_line_indent=Cm(0);cp.paragraph_format.line_spacing=1.0;cp.paragraph_format.space_before=Pt(3);cp.paragraph_format.space_after=Pt(3)
                    cp.alignment=WD_ALIGN_PARAGRAPH.CENTER if rowvalues is t['headers'] or len(str(text))<23 else WD_ALIGN_PARAGRAPH.JUSTIFY
                    for r in cp.runs:r.font.name='Times New Roman';r.font.size=Pt(11);r.bold=rowvalues is t['headers']
                    tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        item=OxmlElement('w:'+edge);item.set(qn('w:w'),'85');item.set(qn('w:type'),'dxa');margins.append(item)
                    tcpr.append(margins)
                    borders=OxmlElement('w:tcBorders')
                    for edge in ['top','left','bottom','right']:
                        item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'5');item.set(qn('w:color'),'D9D9D9');borders.append(item)
                    tcpr.append(borders)
                    if rowvalues is t['headers']:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E2E9EF');tcpr.append(shade)
                trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
                if rowvalues is t['headers']:
                    repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            p(t['source'],small=True)
            markdown.extend([caption,'','| '+' | '.join(t['headers'])+' |','| '+' | '.join(['---']*len(t['headers']))+' |']+['| '+' | '.join(map(str,row))+' |' for row in t['rows']]+['',t['source'],''])
        elif kind=='figure':
            f=FIGURES[key];number=indices['figure'][key]
            path=ROOT/f.get('path',f'outputs/charts/{key}.png')
            para=p('',center=True);para.paragraph_format.keep_with_next=True
            width=8.5 if key=='workbook_kpi' else 12.2 if key=='08_spearman_correlation' else BODY_WIDTH
            run=para.add_run();inline=run.add_picture(str(path),width=Cm(width));inline._inline.docPr.set('descr',f['caption'])
            cap=p(f"Hình {number}: {f['caption']}",'Caption',center=True);cap.paragraph_format.keep_with_next=True;bookmark(cap,f'figure_{number}')
            source_note='Nguồn: thiết kế theo mã và dữ liệu dự án. ' if key in {'architecture','erd'} else 'Nguồn: xử lý từ Mendeley v2 [1], cùng run_id. '
            p(source_note+f['note'],small=True)
            markdown.extend([f"![Hình {number}: {f['caption']}]({path.relative_to(DOCS).as_posix() if path.is_relative_to(DOCS) else '../'+path.relative_to(ROOT).as_posix()})",'',resolve(source_note+f['note']),''])
        elif kind=='equation':
            para=p('',center=True);para.paragraph_format.space_before=Pt(6);para.paragraph_format.space_after=Pt(9)
            math=wrapper('m:oMath',equation_nodes(key));para._p.append(math)
            para.add_run(f"    ({indices['equation'][key]})").font.size=Pt(11)
            readable={'unit_price':'uᵢ = Pᵢ/Aᵢ, Aᵢ > 0','iqr':'IQR = Q₃ − Q₁; L = Q₁ − 1,5 IQR; U = Q₃ + 1,5 IQR','spearman':'ρₛ = corr(R(A), R(P))','bootstrap':'CI = [q*₀,₀₂₅; q*₀,₉₇₅], B = 500','zscore':'zᵢⱼ = (ln(xᵢⱼ) − μⱼ)/σⱼ','kmeans':'J = ∑ⱼ ∑ᵢ∈Cⱼ ‖zᵢ − cⱼ‖²','centroid':'cⱼ = (1/|Cⱼ|) ∑ᵢ∈Cⱼ zᵢ','silhouette':'sᵢ = (bᵢ − aᵢ)/max(aᵢ,bᵢ)'}
            markdown.extend([readable[key]+f"  ({indices['equation'][key]})",''])
        elif kind=='code':
            if key=='sql_search':code=a['sql_benchmark']['query']+';'
            else:
                lines=(ROOT/'src/pipeline.py').read_text(encoding='utf-8').splitlines()
                start=next(i for i,line in enumerate(lines) if 'distances = ((x[:, None, :] - centers' in line)
                code='\n'.join(line[12:] for line in lines[start:start+2])
                code+='\nupdated = np.array([\n    x[labels == j].mean(axis=0) if np.any(labels == j)\n    else x[np.argmax(distances.min(axis=1))]\n    for j in range(k)\n])'
            for line in code.splitlines():
                para=p(line);para.alignment=WD_ALIGN_PARAGRAPH.LEFT;para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.line_spacing=1.0;para.paragraph_format.space_after=Pt(1);para.paragraph_format.keep_with_next=True
                for r in para.runs:r.font.name='Consolas';r.font.size=Pt(9.5)
            para=p('Nguồn: mã thực thi trong src/pipeline.py; chỉ ngắt dòng để trình bày. SQL đầy đủ ở sql/02_queries.sql.',small=True)
            markdown.extend(['```'+('sql' if key=='sql_search' else 'python'),code,'```',''])
    assert chapter_no==6
    para=p('Tài liệu tham khảo','Heading 1');bookmark(para,'references')
    references=[
        '[1] Tran, D. C., Nguyen, H. H., và Ha, H. S. (2021). Vietnam Real Estate Data, Version 2. Mendeley Data. DOI: 10.17632/5477zhnmn2.2. https://data.mendeley.com/datasets/5477zhnmn2/2. CC BY 4.0.',
        '[2] Tài liệu học phần Nhập môn khoa học dữ liệu. Chương 1 Mở đầu. Tệp Chuong 1 - Mo dau.pdf do người học cung cấp; đối chiếu trang PDF 3–5 và 12.',
        '[3] Tài liệu học phần Nhập môn khoa học dữ liệu. Chương 2 Chuẩn bị dữ liệu. Tệp Chuong 2 - Chuan bi du lieu.pdf; các trang về chất lượng, ngoại lệ, Z-score, lấy mẫu và tổng hợp; xem references/DOI_CHIEU_SLIDE.md.',
        '[4] Tài liệu học phần Nhập môn khoa học dữ liệu. Chương 3 Trực quan hóa dữ liệu. Tệp Chuong 3 - Truc quan hoa du lieu.pdf; các trang về mục đích và lựa chọn biểu đồ.',
        '[5] Tài liệu học phần Nhập môn khoa học dữ liệu. Chương 4 Học máy phần 2. Tệp Chương 4 - Hoc may - 2.pdf; nội dung K-means và phân cụm.',
        '[6] Tài liệu học phần Nhập môn khoa học dữ liệu. Chương 5 Cơ sở dữ liệu và SQL phần 1. Tệp Chuong 5 - Co so du lieu va SQL - 1.pdf; mô hình quan hệ, khóa và SQL.',
        '[7] Tài liệu học phần Nhập môn khoa học dữ liệu. Chương 5 Cơ sở dữ liệu và SQL phần 2. Tệp Chuong 5 - Co so du lieu va SQL - 2.pdf; nội dung truy vấn và tổ chức truy cập dữ liệu.',
        '[8] NumPy Developers. NumPy v2.3 Manual, Broadcasting. https://numpy.org/doc/2.3/user/basics.broadcasting.html. Truy cập 16/10/2026.',
        '[9] pandas Developers. pandas 2.3.3 Documentation, Group by split apply combine. https://pandas.pydata.org/pandas-docs/version/2.3/user_guide/groupby.html. Truy cập 16/10/2026.',
        '[10] Matplotlib Development Team. Matplotlib 3.10.6 Documentation, matplotlib.pyplot.hist. https://matplotlib.org/3.10.6/api/_as_gen/matplotlib.pyplot.hist.html. Truy cập 16/10/2026.',
        '[11] SQLite. Query Planning; Foreign Key Support; CREATE TABLE; EXPLAIN QUERY PLAN. https://www.sqlite.org/queryplanner.html; https://www.sqlite.org/foreignkeys.html; https://www.sqlite.org/lang_createtable.html; https://www.sqlite.org/eqp.html. Truy cập 16/10/2026.',
        '[12] Python Software Foundation. Python Standard Library, tài liệu về unicodedata, hashlib, pathlib và urllib.request. https://docs.python.org/3/library/. Truy cập 16/10/2026.',
        '[13] Arthur, D., và Vassilvitskii, S. (2007). k-means++ The Advantages of Careful Seeding. Proceedings of SODA, tr. 1027–1035. https://theory.stanford.edu/~sergei/papers/kMeansPP-soda.pdf.',
        '[14] Rousseeuw, P. J. (1987). Silhouettes A graphical aid to the interpretation and validation of cluster analysis. Journal of Computational and Applied Mathematics, 20, tr. 53–65. DOI: 10.1016/0377-0427(87)90125-7.',
        '[15] Hubert, L., và Arabie, P. (1985). Comparing partitions. Journal of Classification, 2, tr. 193–218. DOI: 10.1007/BF01908075.',
        '[16] Efron, B. (1979). Bootstrap Methods Another Look at the Jackknife. The Annals of Statistics, 7(1), tr. 1–26. DOI: 10.1214/aos/1176344552.',
        '[17] Kho mã dự án KHDL-BatDongSan. https://github.com/vandat-snk/KHDL-BatDongSan. Mã xử lý và bộ bằng chứng mốc 6289bbf; run_id 20261016T015620Z-f2e661a6. Bản báo cáo được bổ sung sau mốc phân tích này.',
    ]
    markdown.extend(['# Tài liệu tham khảo',''])
    for text in references:
        para=p(text);para.paragraph_format.line_spacing=1.0;para.paragraph_format.space_after=Pt(4);para.paragraph_format.keep_together=True
        for run in para.runs:run.font.size=Pt(11)
        markdown.extend([text,''])
    para=p('Phụ lục tái lập và nguồn minh chứng','Heading 1');bookmark(para,'appendix')
    appendix=[
        'Bộ nguồn chotot_data.json được lưu nguyên bản tại data/raw trên máy thực hiện, không đưa lên Git do kích thước và nội dung chi tiết. SHA-256: ef39f153f48a55af07baa86883cd9036652f6d52232fed0be6e8991c7a5bf811. Script tải kiểm đúng bản này trước khi phân tích. Các bản CSV phục vụ công khai không giữ tiêu đề, mô tả, thông tin liên hệ hoặc địa chỉ đường chi tiết.',
        'Nguồn số liệu của báo cáo: outputs/analysis.json, run_manifest.json và các bảng trong outputs/tables. Bằng chứng kiểm: AUDIT_RESULTS.json, ASSISTANT_CHECKS.json, dashboard_test_results.json, workbook_test_results.json và portability_check.json. Hình được lấy từ outputs/charts; hai sơ đồ được dựng từ lược đồ và luồng mã thật. Mọi hình đều kèm số liệu/ví dụ và nguồn trong nội dung chương.',
        'Rà soát nội dung bổ sung được lưu ở docs/DOI_CHIEU_60_TIN_TRO_LY.csv và KET_QUA_KIEM_TRA_NHAN.md. Đây là diễn giải do trợ lý hỗ trợ, tách khỏi các cột nhãn người thật của mẫu gốc. Không sử dụng các ghi chú này để công bố tỷ lệ chính xác nhãn hoặc giả định đã có hai người gán nhãn độc lập.',
        'Để chạy lại phần cốt lõi: tạo môi trường Python 3.11 trở lên, cài requirements.txt, rồi chạy python scripts/run_project.py. Khi dùng đúng raw đã có, thêm --skip-download. Muốn dựng chính xác workbook, môi trường còn cần Node và Artifact Tool; không coi pip install là đã cài phụ thuộc đó. Chạy lại làm đổi run_id và các bằng chứng gắn với lần chạy; báo cáo cần đối soát lại trước khi dùng số liệu mới.',
        'Việc xây dựng mã, rà soát và biên soạn có sử dụng trợ lý AI. Tài liệu không gán các thao tác kiểm bằng công cụ thành xác nhận độc lập của từng thành viên. Những nhận định được giữ theo kết quả lưu và giới hạn kiểm chứng; hồ sơ phân công dùng để tiếp quản nội dung, không thay thế bằng chứng đóng góp đã thực hiện.',
    ]
    markdown.extend(['# Phụ lục tái lập và nguồn minh chứng',''])
    for text in appendix:p(text);markdown.extend([text,''])
    doc.core_properties.title=TITLE;doc.core_properties.subject='Tiểu luận Nhập môn khoa học dữ liệu gồm sáu chương'
    doc.core_properties.author='Võ Văn Đạt; Nguyễn Thị Yến Nhi; Hứa Như Quỳnh'
    doc.core_properties.keywords='bất động sản, giá rao, chất lượng dữ liệu, trực quan hóa, K-means, SQLite'
    doc.save(TARGET)
    (DOCS/'BAO_CAO_TIEU_LUAN_6_CHUONG.md').write_text('\n'.join(markdown),encoding='utf-8')
    log={'generated_at_utc':datetime.now(timezone.utc).isoformat(),'core_run_id':a['run_id'],'document':str(TARGET.relative_to(ROOT)),
         'document_sha256':sha(TARGET),'source_sha256':sha(SOURCE),'figures':figure_count,'tables':table_count,'native_equations':len(indices['equation']),
         'chapter_titles':chapters,'caption_index':indices,'previous_drafts_untouched':True,'render_status':'pending',
         'analysis_sha256':sha(ROOT/'outputs/analysis.json'),
         'figure_assets_sha256':{str(path.relative_to(ROOT)):sha(path) for path in sorted(ASSETS.glob('*.png'))}}
    (DOCS/'report_src/build_manifest.json').write_text(json.dumps(log,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(log,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
