"""Revise the supplied report; preserve its cover and all native equations.

Run with the bundled document Python from the project root. Figures are generated
separately from published numeric results; analytical rules and labels are unchanged.
"""
from pathlib import Path
import csv, hashlib, json, re
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.text.paragraph import Paragraph

ROOT=Path(__file__).resolve().parents[2]
FOLDER=Path(__file__).resolve().parent
SOURCE=Path('E:/UsersData/BAO_CAO_TIEU_LUAN_CAP_NHAT.docx')
DEST=ROOT/'docs/BAO_CAO_TIEU_LUAN_DA_RA_SOAT.docx'
SOURCE_HASH='8fef4131981e12adcf25b77e5b5d4e8b2929e573f8004c50ea7698ced8288e3e'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==SOURCE_HASH
d=Document(SOURCE)
body=d.element.body
original=list(body)
P={i:Paragraph(e,d._body) for i,e in enumerate(original) if e.tag==qn('w:p')}
def setp(i,text):P[i].text=text
def remove(e):
    if e.getparent() is not None:e.getparent().remove(e)
def para(text='',style='Normal'):
    return d.add_paragraph(text,style=style)
def before(anchor,p):anchor.addprevious(p._p);return p
def csvrows(rel):
    with (ROOT/rel).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def vn(x,digits=2):
    return f'{float(x):,.{digits}f}'.replace(',','X').replace('.',',').replace('X','.')

# python-docx omits cached text inside simple SEQ fields from Paragraph.text.
# Preserve the visible caption numbers before replacing navigation and bookmarks.
for p in d.paragraphs:
    if p.style.name=='Caption':
        p.text=''.join(x.text or '' for x in p._p.iter(qn('w:t')))

# Remove stale navigation and the two uncaptioned result blocks after Chapter 6.
for e in original[44:77]:remove(e)
for e in original[300:310]:remove(e)
for i in range(312,318):remove(original[i])
for e in list(body.iter(qn('w:bookmarkStart')))+list(body.iter(qn('w:bookmarkEnd'))):remove(e)

setp(3,'Học phần Nhập môn khoa học dữ liệu')
setp(82,P[82].text.split(' Điều này phù hợp')[0])
setp(91,'Phương pháp thực hiện đi từ kiểm nguồn và chuẩn bị dữ liệu đến phân tích mô tả, kiểm độ nhạy và truyền đạt kết quả. Việc dùng K-means và SQLite có vai trò hỗ trợ cùng mục tiêu khám phá: phân cụm khảo sát cấu trúc của các tin, còn truy vấn quan hệ giúp kiểm và tra cứu kết quả. Mỗi phép phân tích được nối với quy tắc, mã thực thi và đầu ra lưu để có thể đối chiếu thay vì chỉ trình bày khái niệm.')
setp(101,P[101].text.replace(' [3]',''))
setp(111,'Hai đặc trưng phân cụm chính là log tự nhiên của diện tích và đơn giá. Log giảm chênh lệch quy mô; Z-score dùng trung bình và độ lệch chuẩn tổng thể của tập fit để co giãn hai biến. Tuy nhiên, log đơn giá bằng log tổng giá trừ log diện tích, nên diện tích tham gia cả hai đặc trưng. Vì vậy, đề tài đã kiểm bổ sung cặp log diện tích–log tổng giá, thang gốc và chuẩn hóa trung vị/IQR. Chương 5 báo kết quả trên cùng ID, tách riêng độ ổn định khi lấy mẫu với độ nhạy khi đổi hình học của dữ liệu.')
setp(114,P[114].text.replace('[5]','[13]'))
setp(122,P[122].text.replace('[6], [11]','[11]'))
setp(124,P[124].text.replace('[7], [11]','[11]'))
setp(125,P[125].text.replace('[4], [10]','[10]'))
setp(142,'Bảng 3: Từ điển 18 biến của dữ liệu sạch')
setp(170,'Các quy tắc kiểm chéo phát hiện một số mâu thuẫn số học và ngữ cảnh; ngưỡng chênh ít nhất năm lần không bắt mọi khác biệt. Dòng 5093 có 9,7 tỷ ở trường giá và 9,9 tỷ trong mô tả, được giữ kèm cờ source_price_discrepancy_requires_review. Dòng 8151 có 18.000 m² ở trường lõi và 180.000 m² trong mô tả; dòng 8535 có 6 m² ở trường lõi trong khi mô tả nêu 55/60 m². Hai trường hợp sau chưa được quy tắc hiện tại nhận diện đầy đủ. Cả ba còn trong tập EDA; dòng 8535 ngoài tập fit vì cờ IQR. Cờ thống kê không quyết định diện tích nào đúng và không thay việc rà lại nguồn.')
setp(180,'Nguồn: phần gán nhãn và cập nhật tâm của hàm kmeans trong src/pipeline.py; chỉ ngắt dòng để trình bày.')
setp(192,P[192].text.replace('trong main dựng','trong hàm main dựng'))
setp(194,P[194].text+' Riêng scripts/complete_analysis.py kiểm tập ID đánh giá, đối chiếu các trường, tổng hợp tỉnh × loại × khoảng diện tích và chạy 65 phép fit bổ sung. Script kiểm hash các tệp đã khóa trước khi tính; kết quả mở rộng gắn core_run_id, hash script và hash tệp nhãn. Việc sửa nhãn không tự cập nhật các bảng đã xuất.')
setp(235,P[235].text.replace('43/58 trường loại hình xác định khớp','43 trong 58 nhãn loại hình xác định khớp'))
setp(250,P[250].text+' Bảng 20 và Bảng 21 trình bày cả số khớp lẫn độ bao phủ để tránh che khuất những trường chưa kết luận được.')
setp(261,P[261].text.replace('3406','3.406'))
setp(262,P[262].text+' Bảng 23 và Hình 11 tổng hợp cùng phép so sánh ARI; đây là bằng chứng độ nhạy, không phải kết quả kiểm định phân khúc thị trường.')
setp(273,'Đối với dashboard, JavaScript thực thi trên DOM tối giản đã kiểm các trạng thái ở Bảng 18. Tìm Hà Nội trong ô bảng trả 1.014 dòng cho bảng/CSV trong khi KPI toàn trang giữ 6.911 nếu chưa đổi bộ lọc cấp trang. Các kiểm này xác nhận logic lọc, tìm kiếm, phân trang và xuất CSV; chúng chưa xác nhận bố cục hiển thị hoặc mọi tương tác trong trình duyệt thật của bản hiện hành. Bằng chứng sử dụng giao diện của phiên bản cũ không được dùng thay nghiệm thu phiên bản mới.')
setp(282,'Bảng 19 tổng hợp bằng chứng của lần chạy 27/10/2026: dữ liệu sạch 6.911 dòng, tin cách ly, lineage, HTML và Excel được đối soát cùng core_run_id. Đợt rà soát 29/10/2026 chạy lại 45 test Python trên Windows với Python 3.13.7, kiểm logic dashboard và đối soát sản phẩm đã lưu; đồng thời tính lại các bảng mở rộng và 65 phép fit từ code tại commit 549df1a. Kết quả số khớp bộ đã lưu, nhưng thao tác kiểm này không tạo nhãn người độc lập hoặc thay nghiệm thu Microsoft Excel thật.')
setp(287,P[287].text+' Tổng số dòng được bảo toàn trong lineage nên các quyết định giữ, cách ly và gộp trùng có thể kiểm ở cấp bản ghi.')
setp(291,'Đối với câu hỏi 2, tương quan diện tích–tổng giá toàn mẫu gần 0 nhưng đạt khoảng 0,610 ở căn hộ và 0,643 ở nhà tại Hà Nội. Với câu hỏi 3, so sánh trong cùng loại và khoảng diện tích cho kết quả khác phép gộp: nhóm nhà 30–60 m² tại Hà Nội có đơn giá trung vị 83,10 triệu/m² so với 57,25 tại TP.HCM; ở căn hộ cùng khoảng, Hà Nội là 18,18 và TP.HCM là 27,63. Đây là so sánh mô tả trong mẫu, còn khác vị trí và chất lượng chưa được kiểm soát. Với câu hỏi 4, bỏ cờ IQR nhóm làm trung vị giảm 4,62%, trong khi bỏ nhóm có cờ dấu chấm nghìn làm tăng 8,88%; lựa chọn xử lý ảnh hưởng rõ đến tập được mô tả.')
setp(292,'Bảng 24: Mức đáp ứng mục tiêu theo bằng chứng hiện có')
setp(295,P[295].text.replace('Bảng 20','Bảng 24'))
setp(298,'Kết quả 65 phép fit cho thấy cấu trúc cụm tương đối ổn định với các seed và mẫu con đã thử nhưng nhạy với đặc trưng và chuẩn hóa. ARI khoảng 0,0071 khi dùng thang gốc và 0,1833 khi đổi sang log tổng giá không ủng hộ việc gán tên phân khúc cố định. Bước phát triển tiếp theo nên ưu tiên nhãn tham chiếu có xác nhận và ý nghĩa diện tích trước khi thêm mô hình dự báo. Đối với SQL, cần lặp khối đo có thứ tự ngẫu nhiên và đo cả chi phí ghi/dung lượng nếu mục tiêu nghiên cứu chuyển sang đánh giá hiệu năng.')
setp(299,'Về bàn giao, cần duy trì số liệu, code, báo cáo và metadata cùng phiên bản; quy tắc xử lý mới phải có kiểm hồi quy và không dùng lại kết quả nhãn để tự đánh giá cải thiện trên chính mẫu đó. Việc mở HTML hoặc Excel có sẵn khác với khả năng dựng lại toàn bộ trên một máy mới. Các phụ thuộc dựng workbook và quyền sử dụng nguồn cần được ghi trong hướng dẫn triển khai; đóng góp của nhóm và hỗ trợ công cụ được ghi theo công việc thực tế.')
setp(327,'[17] Kho mã dự án KHDL-BatDongSan. Mã phân tích được đối chiếu tại commit 549df1aa0fcb46b011863a02930241807489e91b trên GitHub, nhánh improve-data-validation. https://github.com/vandat-snk/KHDL-BatDongSan/tree/549df1aa0fcb46b011863a02930241807489e91b. Bộ kết quả dùng run_id 20261027T065401Z-5c8df2cc; hồ sơ băm liên kết mã và sản phẩm.')
setp(331,'Mẫu đánh giá nằm trong docs/manual_validation/sample_manifest.json. Bộ nhãn XLSX và lịch sử sửa có văn bản nguồn được giữ riêng trong private, không đưa nội dung liên hệ lên Git. Để tái lập từ tệp nhãn gốc, chạy scripts/complete_analysis.py với --labels trỏ đến tệp được giữ riêng. Bản công khai evaluation_rows.csv cho phép tính lại các bảng đối chiếu từ giá trị tham chiếu đã công bố; thao tác đó không xác minh độc lập người gán nhãn hoặc tái tạo lịch sử sửa trong XLSX.')
setp(333,'Đợt 29/10/2026 đối soát lại dữ liệu đã lưu, chạy 45 test Python, kiểm logic JavaScript và tính lại 65 phép fit cùng các bảng mở rộng trên Windows. Việc tái tính giữ nguyên quy tắc và nhãn tham chiếu công khai. Hash khóa quy tắc được kiểm lại sau khi bảo toàn byte xuống dòng qua Git; không sửa hash để chấp nhận thay đổi nội dung. Lần phân tích chính 27/10/2026 vẫn là bộ số liệu được trình bày, còn lượt tái kiểm là bằng chứng đối chiếu bổ sung.')

# Replace course-page citations by primary technical references.
mapping={1:1,8:2,9:3,10:4,11:5,12:6,13:7,14:8,15:9,16:10,17:11}
for p in list(d.paragraphs):
    if p._p.getparent() is None or p._p.xpath('.//m:oMath') or p._p.xpath('.//w:drawing'):continue
    if re.search(r'\[\d+\]',p.text):
        def sub(match):
            old=int(match[1]);assert old in mapping,(old,p.text)
            return f'[{mapping[old]}]'
        p.text=re.sub(r'\[(\d+)\]',sub,p.text)

references={
311:'[1] Tran, D. C., Nguyen, H. H., và Ha, H. S. (2021). Vietnam Real Estate Data, Version 2. Mendeley Data. DOI 10.17632/5477zhnmn2.2. https://data.mendeley.com/datasets/5477zhnmn2/2. Giấy phép CC BY 4.0.',
318:'[2] NumPy Developers. NumPy v2.3 Manual. Broadcasting. https://numpy.org/doc/2.3/user/basics.broadcasting.html. Truy cập 29/10/2026.',
319:'[3] pandas Developers. pandas 2.3.3 Documentation. Group by split apply combine. https://pandas.pydata.org/pandas-docs/version/2.3/user_guide/groupby.html. Truy cập 29/10/2026.',
320:'[4] Matplotlib Development Team. Matplotlib 3.10.6 Documentation. matplotlib.pyplot.hist. https://matplotlib.org/3.10.6/api/_as_gen/matplotlib.pyplot.hist.html. Truy cập 29/10/2026.',
321:'[5] SQLite. Query Planning; Foreign Key Support; CREATE TABLE; EXPLAIN QUERY PLAN. https://www.sqlite.org/queryplanner.html; https://www.sqlite.org/foreignkeys.html; https://www.sqlite.org/lang_createtable.html; https://www.sqlite.org/eqp.html. Truy cập 29/10/2026.',
322:'[6] Python Software Foundation. Python 3.12 Standard Library. Các mô-đun unicodedata, hashlib, pathlib và urllib.request. https://docs.python.org/3.12/library/. Truy cập 29/10/2026.',
323:'[7] Arthur, D., và Vassilvitskii, S. (2007). k-means++: The Advantages of Careful Seeding. Proceedings of SODA, 1027–1035. https://theory.stanford.edu/~sergei/papers/kMeansPP-soda.pdf.',
324:'[8] Rousseeuw, P. J. (1987). Silhouettes: A graphical aid to the interpretation and validation of cluster analysis. Journal of Computational and Applied Mathematics, 20, 53–65. DOI 10.1016/0377-0427(87)90125-7.',
325:'[9] Hubert, L., và Arabie, P. (1985). Comparing partitions. Journal of Classification, 2, 193–218. DOI 10.1007/BF01908075.',
326:'[10] Efron, B. (1979). Bootstrap Methods: Another Look at the Jackknife. The Annals of Statistics, 7(1), 1–26. DOI 10.1214/aos/1176344552.'}
for i,text in references.items():setp(i,text)

# State the evaluation design before the design/implementation chapters.
anchor=original[129]
before(anchor,para('Đánh giá quy tắc tách mẫu phát triển khỏi mẫu đối chiếu. Từ các bản ghi không trùng, thiết kế chọn 180 ID theo ba nhóm của baseline: tin giữ có giao dịch giả định, tin giữ còn lại và tin cách ly. Trong mỗi nhóm, loại suy luận và mức rủi ro tạo các tầng lấy mẫu cân bằng; 120 ID phục vụ phát triển và 60 ID phục vụ đối chiếu, mỗi nhóm baseline có 20 ID đối chiếu. Các ID đã biết trong lượt rà soát trước được loại khỏi quá trình chọn. Mẫu cân bằng có xác suất chọn khác nhau giữa các tầng nên tỷ lệ gộp không phải ước lượng độ đúng của toàn nguồn. Tin gần trùng khác hash vẫn có thể xuất hiện ở hai tập.'))
before(anchor,para('Mỗi trường chỉ được chấm khi tham chiếu đủ xác định. Số khớp chia số tham chiếu xác định mô tả mức khớp có điều kiện; số tham chiếu xác định chia 60 mô tả độ bao phủ. Unknown, số trống và quyết định review được báo riêng. Nhãn máy other_unknown hoặc sale_assumed được coi là chưa khẳng định, không tính khớp với một nhãn cụ thể. So sánh số dùng sai số tương đối 10⁻⁶, sai số tuyệt đối 1 VND cho giá và 10⁻⁶ m² cho diện tích. Việc kiểm giá trị parser của một tin đã cách ly không đồng nghĩa quyết định giữ tin đó đúng.'))
before(anchor,para('Bộ tham chiếu có một người gán nhãn ban đầu và 17 tin được AI hiệu chỉnh; chưa có xác nhận độc lập cho mọi sửa đổi. Vì thế, kết quả chỉ là đối chiếu với bộ tham chiếu hiện có, không gọi là kiểm định nhãn người độc lập và không tính hệ số đồng thuận giữa hai người. Nếu dùng lỗi của tập đối chiếu để sửa quy tắc, cần một tập mới trước khi công bố mức cải thiện.'))

# Extend the module map with the implemented supplementary analysis.
module_table=d.tables[4]
cells=module_table.add_row().cells
for cell,text in zip(cells,['complete_analysis.py','ID khóa, nhãn tham chiếu và dữ liệu cùng run_id','Đối chiếu từng trường, nhóm theo ba khóa, 65 phép fit và metadata']):cell.text=text

# Four captioned result tables and two figures, all within Chapter 5.
anchor=original[286]
def insertp(text,style='Normal'):return before(anchor,para(text,style))
def table(caption,headers,rows):
    insertp(caption,'Caption')
    t=d.add_table(rows=1,cols=len(headers))
    for c,s in zip(t.rows[0].cells,headers):c.text=str(s)
    for values in rows:
        for c,s in zip(t.add_row().cells,values):c.text=str(s)
    anchor.addprevious(t._tbl)
    return t
def figure(caption,path,source,width=15.8):
    p=insertp('');p.add_run().add_picture(str(path),width=Cm(width))
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    insertp(caption,'Caption');insertp(source)

insertp('Kết quả đối chiếu nhãn cần đọc đồng thời mức khớp và độ bao phủ ở Bảng 20. Chẳng hạn, tổng giá đạt 47/50 trên những tham chiếu có số, còn quyết định giữ/cách ly chỉ có 8 tham chiếu đã xác định. Có 52 trường hợp review, nên tỷ số 8/8 chưa đủ kiểm quyết định giữ trên đa số mẫu. Cách ghi cả hai mẫu số giúp tránh một tỷ lệ cao che khuất sự thiếu hụt bằng chứng.')
metrics=csvrows('outputs/extended/evaluation_metrics.csv')
labels={'property':'Loại tài sản','transaction':'Giao dịch','basis':'Cơ sở giá','price':'Tổng giá','area':'Diện tích','status':'Giữ/cách ly'}
table('Bảng 20: Mức khớp và độ bao phủ trên 60 nhãn tham chiếu',
      ['Trường','Khớp/xác định','Chưa xác định','Bao phủ %'],
      [[labels[r['field']],f"{r['n_match']}/{r['n_resolved']}",r['n_unresolved'],vn(float(r['resolved_coverage'])*100)] for r in metrics if r['group']=='all_sample'])
insertp('Nguồn: evaluation_metrics.csv, cùng mốc phân tích. Bao phủ dùng mẫu số 60; mức khớp chỉ tính trên tham chiếu xác định. Tham chiếu có AI hỗ trợ, không phải nhãn độc lập.')
insertp('Bảng 21 giữ riêng ba nhóm được xác định lúc lấy mẫu. Nhóm giao dịch giả định có 0/11 khớp ở trường giao dịch vì máy chọn chưa khẳng định, trong khi tham chiếu đưa ra nhãn cụ thể. Đây không đồng nghĩa cả 11 tin đều là tin thuê; nó cho thấy quy tắc không cung cấp được khẳng định giao dịch trên những tham chiếu này. Việc các tầng có tỷ lệ khác nhau cũng giải thích vì sao không dùng tỷ lệ gộp của mẫu cân bằng để suy rộng toàn nguồn.')
groups={'kept_assumed':'Giữ, giả định bán','kept_other':'Giữ, còn lại','rejected':'Cách ly'}
lookup={(r['group'],r['field']):r for r in metrics}
table('Bảng 21: Đối chiếu nhãn theo ba nhóm lấy mẫu của baseline',
      ['Nhóm baseline','n','Loại tài sản','Giao dịch','Giữ/cách ly'],
      [[name,20]+[f"{lookup[(key,f)]['n_match']}/{lookup[(key,f)]['n_resolved']}" for f in ['property','transaction','status']] for key,name in groups.items()])
insertp('Nguồn: evaluation_metrics.csv. Ba cột cuối ghi số khớp/số tham chiếu xác định của từng trường; phần chưa xác định không tính thành đúng hoặc sai.')
insertp('Để làm rõ ảnh hưởng cơ cấu, Bảng 22 và Hình 10 chọn ba cặp có cùng nhãn loại và khoảng diện tích, với n≥30 ở cả Hà Nội và TP.HCM. Sáu ô được đối chiếu theo nhãn và số mẫu của từng cặp. Trong nhóm nhà 30–60 m², Hà Nội có trung vị cao hơn TP.HCM; ở nhóm căn hộ cùng khoảng diện tích, thứ tự đổi lại. So sánh này cho thấy câu khẳng định một thành phố luôn đắt hơn sẽ vượt bằng chứng; chưa kiểm soát quận, chất lượng, vị trí hoặc định nghĩa diện tích.')
conditional=csvrows('docs/review_20261029/assets/conditional_comparison.csv')
types={'house':'Nhà','apartment':'Căn hộ'}
bands={'(30,60]':'30 < A ≤ 60','(60,100]':'60 < A ≤ 100'}
table('Bảng 22: So sánh đơn giá trong cùng loại và khoảng diện tích',
      ['Địa phương/loại','Khoảng m²','n','Trung vị','Q1–Q3'],
      [[('TP.HCM' if r['province']=='Tp Hồ Chí Minh' else 'Hà Nội')+' / '+types[r['property_type']],bands[r['area_band']],vn(r['n'],0),vn(float(r['median_price_per_m2_vnd'])/1e6),vn(float(r['q1_price_per_m2_vnd'])/1e6)+'–'+vn(float(r['q3_price_per_m2_vnd'])/1e6)] for r in conditional])
insertp('Nguồn: province_type_area_summary.csv; các cột trung vị và Q1–Q3 có đơn vị triệu VND/m². Sáu ô nằm trong 41 nhóm biết tỉnh/loại có n≥30; tổng toàn bộ 41 nhóm là 4.594 tin.')
figure('Hình 10: Đơn giá theo địa phương trong cùng loại và khoảng diện tích',FOLDER/'assets/conditional_comparison.png',
       'Nguồn: số liệu Bảng 22. Điểm là trung vị, đoạn ngang là Q1–Q3, không phải khoảng tin cậy; n từ 38 đến 1.054. Các loại hình là nhãn suy luận.')
insertp('Độ nhạy phân cụm được đánh giá tại cùng k=2 và seed 42, chỉ so ARI trên các ID chung. Bảng 23 và Hình 11 dùng cấu hình log diện tích–log đơn giá, Z-score, bỏ IQR làm mốc. Khi thêm ngoại lệ, mô hình được fit trên 6.911 tin nhưng ARI vẫn so trên 6.652 ID chung. Những cấu hình khác giữ 6.652 tin. Việc cố định k làm rõ ảnh hưởng biểu diễn; không ngầm chọn lại số cụm thuận lợi cho từng so sánh.')
ari=csvrows('outputs/extended/cluster_feature_sensitivity.csv')
scenario_names=['Mốc log + Z-score, bỏ IQR','Giữ ngoại lệ','Thang gốc + Z-score','Log + trung vị/IQR','Log diện tích + log tổng giá']
table('Bảng 23: Độ nhạy của phân hoạch khi thay đổi đặc trưng và chuẩn hóa',
      ['Cấu hình so với mốc','ID chung','ARI'],
      [[name,vn(r['n_common'],0),vn(r['ari_on_common_rows'],4)] for name,r in zip(scenario_names,ari)])
insertp('Nguồn: cluster_feature_sensitivity.csv; cùng k=2 và seed 42. ARI đo sự giống nhau giữa hai phân hoạch, không đo độ đúng so với một phân khúc thị trường đã xác minh.')
figure('Hình 11: Mức thay đổi phân hoạch theo cách biểu diễn dữ liệu',FOLDER/'assets/cluster_feature_ari.png',
       'Nguồn: toàn bộ năm giá trị ở Bảng 23; mỗi so sánh có 6.652 ID chung. ARI lần lượt 1,0000; 0,9764; 0,0071; 0,7221 và 0,1833.')
insertp('Trong 60 phép fit theo cấu hình/k/seed, ARI nhỏ nhất giữa các seed đã thử là 0,9778. Năm lần fit mẫu con 80% gồm 5.322 tin mỗi lần có ARI 0,9792–0,9980 so với mô hình mốc trên ID chung. Các kết quả này hỗ trợ tính lặp lại trong những phép thử đã làm; số lần thử không được coi là số nghiên cứu độc lập. Điểm silhouette không xác định khi mẫu chỉ chứa một cụm được để trống và ghi lý do. Điểm ở các không gian biến đổi khác nhau chỉ dùng thảo luận hình học, không làm thước đo tuyệt đối để chọn biến đổi tốt nhất.')
insertp('Lượt tái kiểm 29/10/2026 tính lại các số liệu trên từ quy tắc không thay đổi. Đối chiếu 60 dòng dùng các giá trị tham chiếu công khai; các bảng nhóm và 65 phép fit được chạy lại bằng mã gốc. Các giá trị số khớp trong sai số số thực đã ghi, nhưng không bổ sung bằng chứng xác minh nội dung kinh tế của từng tin.')

# Remove empty spacer paragraphs without removing section breaks, math or pictures.
inside=False
for p in list(d.paragraphs):
    if p._p is original[13]:inside=True
    if not inside:continue
    has_content=bool(p.text.strip() or p._p.xpath('.//w:drawing|.//m:oMath|.//w:sectPr|.//w:br'))
    if not has_content:remove(p._p)
P[37].paragraph_format.page_break_before=True

# Apply restrained academic formatting; the supplied cover remains authoritative.
for style_name in ['Normal','Caption','Heading 1','Title']:
    style=d.styles[style_name]
    style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    for elem in list(style.element.iter(qn('w:pBdr'))):remove(elem)
d.styles['Normal'].font.size=Pt(12)
d.styles['Normal'].paragraph_format.line_spacing=1.15
d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Heading 1'].font.size=Pt(14)
d.styles['Heading 1'].font.bold=True
d.styles['Caption'].font.size=Pt(11)
d.styles['Caption'].font.bold=True
inside=False;references_started=False
for p in d.paragraphs:
    if p._p is original[13]:inside=True
    if not inside:continue
    pf=p.paragraph_format
    if not p.text and p._p.xpath('.//w:sectPr'):
        p.style='Normal';pf.page_break_before=False;pf.keep_with_next=False
        continue
    if p.text in ['LỜI CẢM ƠN','TÓM TẮT']:
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.first_line_indent=Cm(0)
        pf.space_after=Pt(12);pf.keep_with_next=True
        for r in p.runs:r.bold=True;r.font.size=Pt(14)
        continue
    if p.style.name=='Heading 1' and p.text:
        pf.page_break_before=True;pf.keep_with_next=True;pf.first_line_indent=Cm(0)
        pf.space_before=Pt(0);pf.space_after=Pt(12)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if p.text=='Tài liệu tham khảo':references_started=True
        if p.text.startswith('Phụ lục tái lập'):references_started=False
    elif p.style.name=='Caption':
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Cm(0)
        pf.keep_with_next=True;pf.space_before=Pt(5);pf.space_after=Pt(6)
    elif p._p.xpath('.//m:oMath'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_with_next=True
    elif p._p.xpath('.//w:drawing'):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.keep_with_next=True;pf.first_line_indent=Cm(0)
    elif p.text.startswith('Nguồn:'):
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.first_line_indent=Cm(0)
        pf.space_after=Pt(6);pf.keep_with_next=False
        for r in p.runs:r.font.size=Pt(10);r.bold=False
    elif p.text.startswith(('distances =','labels =','updated =','    x[','    else ','    for ','])','SELECT ','FROM ','WHERE ','  AND ','ORDER BY ')):
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.first_line_indent=Cm(0)
        pf.line_spacing=1;pf.space_after=Pt(0);pf.keep_with_next=True
        for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)
    elif p.text:
        p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.first_line_indent=Cm(.7)
        pf.keep_with_next=False
        for r in p.runs:r.bold=False
        if references_started:
            pf.first_line_indent=Cm(0);pf.line_spacing=1.05;pf.space_after=Pt(5)
            for r in p.runs:r.font.size=Pt(10.5)

for t in d.tables:
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    pr=t._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is None:borders=OxmlElement('w:tblBorders');pr.append(borders)
    for edge in ['top','left','bottom','right','insideH','insideV']:
        item=borders.find(qn('w:'+edge))
        if item is None:item=OxmlElement('w:'+edge);borders.append(item)
        for key,val in [('val','single'),('sz','4'),('color','D9D9D9')]:item.set(qn('w:'+key),val)
    for i,row in enumerate(t.rows):
        row.height=None
        rp=row._tr.get_or_add_trPr()
        if i==0 and rp.find(qn('w:tblHeader')) is None:rp.append(OxmlElement('w:tblHeader'))
        if rp.find(qn('w:cantSplit')) is None:rp.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr();shd=cp.find(qn('w:shd'))
            if shd is None:shd=OxmlElement('w:shd');cp.append(shd)
            shd.set(qn('w:fill'),'E2E9EF' if i==0 else 'FFFFFF')
            margins=cp.find(qn('w:tcMar'))
            if margins is None:margins=OxmlElement('w:tcMar');cp.append(margins)
            for side in ['top','bottom','left','right']:
                edge=margins.find(qn('w:'+side))
                if edge is None:edge=OxmlElement('w:'+side);margins.append(edge)
                edge.set(qn('w:w'),'70' if side in ['top','bottom'] else '85');edge.set(qn('w:type'),'dxa')
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
                p.paragraph_format.keep_with_next=False
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or re.fullmatch(r'[\d.,%/–− -]+',p.text) else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(10.5);r.bold=(i==0);r.font.color.rgb=RGBColor(0,0,0)

# Rebuild chapter navigation and separate figure/table lists using live PAGEREF.
bookmark_counter=2000
def bookmark(p,name):
    global bookmark_counter
    bookmark_counter+=1
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_counter));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_counter))
    p._p.insert(1 if p._p.pPr is not None else 0,start);p._p.append(end)
chapters=[];figures=[];tables=[]
for p in d.paragraphs:
    if p.style.name=='Heading 1':
        match=re.match(r'Chương (\d)',p.text)
        name='chapter_'+match[1] if match else ('references' if p.text=='Tài liệu tham khảo' else 'appendix')
        bookmark(p,name);chapters.append((p.text,name))
    elif p.style.name=='Caption':
        match=re.match(r'(Hình|Bảng) (\d+):',p.text)
        assert match,p.text
        name=('figure_' if match[1]=='Hình' else 'table_')+match[2]
        bookmark(p,name)
        (figures if match[1]=='Hình' else tables).append((p.text,name))
assert len(chapters)==8 and len(figures)==11 and len(tables)==24,(chapters,len(figures),len(tables),figures,tables)
nav_anchor=original[77]
def navp(text):return before(nav_anchor,para(text))
def navtitle(text):
    p=navp(text);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(12)
    for r in p.runs:r.bold=True;r.font.size=Pt(14)
def entry(text,name,size=11):
    p=navp(text+'\t');pf=p.paragraph_format
    pf.first_line_indent=Cm(0);pf.space_after=Pt(4);pf.line_spacing=1.05
    pf.tab_stops.add_tab_stop(Cm(16.0),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    for r in p.runs:r.font.size=Pt(size)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGEREF '+name+' \\h ')
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='0';r.append(t);field.append(r);p._p.append(field)
def new_nav_page():navp('').add_run().add_break(WD_BREAK.PAGE)
navtitle('Mục lục')
for text,name in chapters:entry(text,name,11.5)
new_nav_page();navtitle('Danh mục hình')
for text,name in figures:entry(text,name,10.5)
new_nav_page();navtitle('Danh mục bảng')
for text,name in tables:entry(text,name,10.5)

settings=d.settings.element
update=settings.find(qn('w:updateFields'))
if update is None:update=OxmlElement('w:updateFields');settings.append(update)
update.set(qn('w:val'),'true')
d.core_properties.title='Phân tích giá rao bất động sản Việt Nam'
d.core_properties.subject='Tiểu luận Nhập môn khoa học dữ liệu'
d.core_properties.author='Võ Văn Đạt; Nguyễn Thị Yến Nhi; Hứa Như Quỳnh'
text='\n'.join(p.text for p in d.paragraphs)
for forbidden in ['slide','giáo trình','trang PDF','DOI_CHIEU_SLIDE','chưa khẳng định đã đẩy','chưa phải thí nghiệm đã thực hiện']:
    assert forbidden.lower() not in text.lower(),forbidden
assert len(d.element.xpath('//m:oMath'))==8
assert len(d.tables)==24
DEST.parent.mkdir(exist_ok=True)
d.save(DEST)

# Readable source export for provenance; equations remain native in the DOCX.
lines=[]
for e in d.element.body:
    extract=lambda el:''.join(x.text or '' for x in el.iter() if x.tag in (qn('w:t'),qn('m:t')))
    if e.tag==qn('w:p'):
        p=Paragraph(e,d._body)
        lines.append(('## ' if p.style.name=='Heading 1' else '')+extract(e))
    elif e.tag==qn('w:tbl'):
        lines.extend(' | '.join(extract(c) for c in row.findall(qn('w:tc'))) for row in e.findall(qn('w:tr')))
DEST.with_suffix('.md').write_text('\n\n'.join(lines),encoding='utf-8')
(FOLDER/'build.json').write_text(json.dumps({'source_sha256':SOURCE_HASH,'source_path':str(SOURCE),
    'analysis_commit':'549df1aa0fcb46b011863a02930241807489e91b',
    'core_run_id':'20261027T065401Z-5c8df2cc','figures':11,'tables':24,'native_equations':8,
    'changes':['Removed course-page citations; retained primary sources.',
               'Integrated supplementary results into Chapter 5.',
               'Added field coverage, sampling-stratum results and matched comparisons.',
               'Corrected descriptions of implemented flags, module and feature sensitivity.',
               'Rebuilt all navigation and captions; preserved cover and native math.']},ensure_ascii=False,indent=2),encoding='utf-8')
print(DEST)
